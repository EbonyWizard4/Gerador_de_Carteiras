"""
Cria o Relatório em PDF
"""

import pandas as pd

from docx import Document
from fpdf import FPDF

WIDTH = 210
HEIGHT = 297


class PDF(FPDF):
    """
    Expecificações e configurações do PDF

    Args:
        FPDF (Lib): Permite editar o documento PDF
    """

    def header(self):
        # Logo
        self.image("img/EbwInvest.png", 0, 0, WIDTH)
        self.ln(50)

    # Page footer
    def footer(self):
        # Position at 1.5 cm from bottom
        self.set_y(-15)
        # Arial italic 8
        self.set_font("Arial", "I", 8)
        # Page number
        self.cell(0, 10, "Page " + str(self.page_no()) + "/{nb}", 0, 0, "C")


def do_pdf(destino) :
    """Edição do PDF

    Args:
        Destino (str): Diretório de salvamento do PDF
    """

    # Propriedades do PDF
    pdf = PDF()
    pdf.alias_nb_pages()

    

    # Lista de tabelas
    tab_list = {
            "Magic Formula": "CSV/magicForm.csv",
            "Benjamin Graham": "CSV/ben_grahan.csv",
            "Décio Bazin": "CSV/decio_bazin.csv",
        }

    # Elementos do PDF
    for tab_item, caminho in tab_list.items():
        
        # - Configurações da página
        pdf.add_page()
        titulo = f"As Melhores Ações \nCom Melhor Custo/Benefício.\nSegundo: Método {tab_item}"
        pdf.set_font("Arial", "B", 24)
        pdf.write(10, titulo )
        pdf.ln(40)
                
        # - tabela
        pdf.set_font("Times", "", 12)
        table_data = pd.read_csv(caminho)
        table_data = table_data.applymap(str)
        columns = [list(table_data)]
        rows = table_data.values.tolist()
        data = columns + rows

        with pdf.table(
            borders_layout="MINIMAL",
            cell_fill_color=200,
            cell_fill_mode="ROWS",
            line_height=pdf.font_size * 2.5,
            text_align="CENTER",
        ) as table:
            for data_row in data:
                row = table.row()
                for datum in data_row:
                    row.cell(datum)

    # - Ultima página
    pdf.add_page()
    pdf.set_margin(40)
    pdf.set_font("Arial", "B", 24)
    pdf.write(10, "DISCLAIMER")
    pdf.ln(15)
    pdf.set_margin(20)
    pdf.set_font("Times", "", 12)
    documento = Document("documents/Disclaimer.docx")
    for paragrafo in documento.paragraphs:
        pdf.write(5, paragrafo.text)
        pdf.ln(10)

    pdf.output(destino + "/Carteira.pdf", "F")
